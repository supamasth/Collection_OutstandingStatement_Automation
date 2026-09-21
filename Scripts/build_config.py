# Builds Data\Config.xlsx for Collection_OutstandingStatement_Automation
# Sheet structure per MyRPA root CLAUDE.md "Config.xlsx Design Standards (DKSH BST Style)"
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

HEADER_FILL = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GROUP_FONT = Font(bold=True)

def style_header(ws, cols=3):
    for c in range(1, cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

# ---------------------------------------------------------------------------
# Sheet 1: Settings
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "Settings"
ws.append(["Key", "Value", "Description"])
style_header(ws)

rows = []

def group(title):
    rows.append(["", "", ""])
    rows.append([f"### {title}", "", ""])

def kv(key, value, desc):
    rows.append([key, value, desc])

# Group 1: Process Control
group("Process Control")
kv("RunRound", "1", "Reserved for future multi-round runs. Currently always 1 (single pass over unread emails).")
kv("MaxRetryNumber", "3", "REFramework ApplicationException retry count per transaction.")
kv("RetryDelaySeconds", "5", "Delay between ApplicationException retries.")

# Group 2: Credentials/URLs
group("Credentials / SAP Connection")
kv("SapConnection", "01A  P00 HERMES [PROD]", "Connection name inside SAP Logon.")
kv("SapClientCode", "300", "SAP client number.")
kv("SapApplication", r"C:\Program Files\SAP\FrontEnd\SAPGUI\saplogon.exe", "Full path to saplogon.exe.")
kv("SapCredentialName", "TH HEC SAP RPA S4 SANDBOX", "Orchestrator asset name passed to Get Robot Credential inside Workflows\\OpenSAP.xaml. Must match the Assets sheet row below.")
kv("SapTimeout", "30", "SAP login timeout in seconds.")
kv("SapTCode_Report1", "YFI_OC_GEN_I039", "T-Code for Customer Statement (Report 1).")
kv("SapTCode_Report2", "ZTHSD_OC_GEN_R023", "T-Code for Daily Sales Report, brand-filtered (Report 2).")
kv("SAP_Report1_Variant", "RPAOUTSTANDING", "Variant name executed on YFI_OC_GEN_I039.")
kv("SAP_Report2_SalesOrg", "TH54", "Sales Organization entered on ZTHSD_OC_GEN_R023.")

# Group 3: File Paths
group("File Paths")
kv("OutputBasePath", "W:\\BKK_HEC-COLLECTION\\Outstanding_Statement_Auto\\", "Base folder for per-transaction subfolders (Report 1 + Report 2 outputs).")
kv("DKSHCustomerCsvPath", r"Data\Input\DKSH_Customer.csv", "Customer master file (CustCode, Status, CustGroup_Description).")
kv("MasterExcelPath", r"Data\Input\MasterExcel_ClientAccessDB.xlsx", "Sender authorization list (Client Email, MatGroup, MatGroup4, Status).")
kv("LogFilePath", "Logs\\", "Local log folder.")
kv("ScreenshotFolder", "Screenshots\\", "Folder for exception screenshots.")

# Group 4: Business Rules
group("Business Rules")
kv("CustomerIDPattern", "(?<!\\d)17\\d{7}(?!\\d)", "Regex.Match search pattern (not anchored): finds a standalone 9-digit run starting with 17 inside subject/body text.")
kv("CustGroupExclusionList", "EDUCATIONAL INST.|GOV. HOSPITAL|INTER-COMPANY|MILITARY HOSP. BID|PRIVATE HOSPITAL|THIRD PARTY TENDER", "Pipe-delimited CustGroup_Description exclusion list (Section 3.2). Edit without redeploy.")
kv("SubfolderRetentionDays", "", "TBD — confirm with business whether per-transaction subfolders under OutputBasePath are deleted after send or retained N days. Leave blank = retain indefinitely until decided.")

# Group 5: Outlook / Notifications
group("Outlook Folders")
kv("MailAccount", "rpasupport3@dksh.com", "Mailbox monitored by the robot. NOTE: brief says this changes to telecollection.hec.bkk@dksh.com at go-live — change this one value only.")
kv("OutlookFolder_Intake", "OutstandingRequest", "Folder the robot reads UNREAD items from (Outlook rule moves matching mail here from Inbox).")
kv("OutlookFolder_Processed", "OutstandingRequest\\Processed", "Destination folder after a successful send.")
kv("OutlookFolder_Exception", "OutstandingRequest\\Exception", "Destination folder after any Section 3 validation failure.")
kv("OutlookFolder_ForwardedToCollection", "OutstandingRequest\\ForwardedToCollection", "Destination folder for replies/forwards to the robot's own emails (not a fresh request) - forwarded to the Collection team untouched, not auto-replied to. Create this folder manually in Outlook, same as Processed/Exception.")

group("Notifications")
kv("Mail_Success_To_Extra", "CreditControl2.HEC.BKK@dksh.com;TeleCollection.HEC.BKK@dksh.com;Payin.HEC.BKK@dksh.com", "Additional To recipients for success reply, besides the requester. Placeholder addresses — confirm exact mailbox names with Collection team.")
kv("Mail_Success_CC", "Chitri.Lugsaniyanont@dksh.com", "CC on success reply. Placeholder — confirm exact mailbox.")
kv("Mail_Exception_CC", "telecollection.hec.bkk@dksh.com", "CC on every Section 3 exception reply (Section 6.2 — no exceptions to this rule).")
kv("RPAAdminEmail", "", "To address for System Exception emails (Section 6.3). Fill in before go-live — not sent to requester or Collection team.")

for r in rows:
    ws.append(r)

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=1):
    cell = row[0]
    if isinstance(cell.value, str) and cell.value.startswith("###"):
        cell.font = GROUP_FONT

widths = {"A": 32, "B": 55, "C": 70}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# ---------------------------------------------------------------------------
# Sheet 2: Assets
# ---------------------------------------------------------------------------
ws2 = wb.create_sheet("Assets")
ws2.append(["Key", "OrchestratorName", "Description"])
style_header(ws2)
ws2.append(["CredentialSAP", "TH HEC SAP RPA S4 SANDBOX", "SAP logon credential (username/password) fetched from Orchestrator Asset at Init. Must match Settings!SapCredentialName."])
ws2.column_dimensions["A"].width = 28
ws2.column_dimensions["B"].width = 42
ws2.column_dimensions["C"].width = 70

wb.save(r"C:\Users\hecrpasupport.th\Documents\MyRPA\Collection_OutstandingStatement_Automation\Data\Config.xlsx")
print("Config.xlsx written.")
